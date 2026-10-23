from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT_PATH = r"D:\CodexProject\2026年校电赛\iq_modulation_summary_am_verified.docx"
PROJECT_PATH = (
    r"D:\VSCodeCMAKEProject\Simple wireless automatic transceiver device"
    r"\Simple_wireless_automatic_transceiver_deviceV2.0"
    r"\Core\APP\DAC\app_dac_wavegen.c"
)


def set_page(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)


def set_font(run, size_pt: float, bold: bool = False, color=None) -> None:
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size_pt)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_text(doc: Document, text: str, style: str = "Normal"):
    p = doc.add_paragraph(style=style)
    if style == "Normal":
        p.paragraph_format.line_spacing = 1.08
        p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    if style == "Title":
        set_font(run, 20, True)
    elif style.startswith("Heading"):
        set_font(run, 14 if style == "Heading 1" else 12, True, RGBColor(31, 78, 121))
    else:
        set_font(run, 11)
    return p


def add_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    set_font(run, 11)


def add_footer_with_page_number(doc: Document) -> None:
    section = doc.sections[0]
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_PARAGRAPH_ALIGNMENT.RIGHT
    run = footer.add_run("Page ")
    set_font(run, 9)

    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")

    footer._p.append(fld_begin)
    footer._p.append(instr)
    footer._p.append(fld_end)


def build_doc() -> None:
    doc = Document()
    set_page(doc)
    add_footer_with_page_number(doc)

    title = doc.add_paragraph()
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    run = title.add_run("IQ 调制公式与 AM 实现核实总结")
    set_font(run, 22, True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    run = subtitle.add_run("针对 app_dac_wavegen.c 当前实现")
    set_font(run, 11, False, RGBColor(102, 102, 102))
    subtitle.paragraph_format.space_after = Pt(12)

    add_text(doc, "1. AM 实现是否正确", "Heading 1")
    add_bullet(
        doc,
        "有依据：按当前代码本身的注释与数学形式看，它不是标准“带载波 AM”的完整基带表达。"
        "代码位置在 app_dac_wavegen.c 第 423~431 行。",
    )
    add_bullet(
        doc,
        "当前实现公式为：I_signed = A * m * sin(wm t)，Q_signed = 0；"
        "随后再统一加 DAC 直流偏置 code。"
        "这相当于只输出调幅分量，没有在算法层显式加入常数载波项。",
    )
    add_bullet(
        doc,
        "如果目标是标准 AM 的复包络，通常应写成：I = A * (1 + m * sin(wm t))，Q = 0。"
        "也就是在 I 路里同时包含常数项和调制项。",
    )
    add_bullet(
        doc,
        "因此，单从数字公式看，当前 AM 更接近“只保留调制分量”的实现；"
        "是否最终还能在射频侧形成带载波 AM，要取决于后级模拟链路是否把 DAC 偏置当成有效载波项。"
        "这一点仅靠当前代码不能下定论。",
    )

    add_text(doc, "2. 代码中的公共换算", "Heading 1")
    add_bullet(
        doc,
        "相位步进：phase_step = freq_hz * 2^32 / Fs，代码在 app_dac_wavegen.c 第 377~380 行。"
        "其中采样率 Fs = 2 MHz。",
    )
    add_bullet(
        doc,
        "幅度换算：peak_code = mv_to_code(vpp_mv / 2)，代码在第 361~375 行。"
        "也就是配置的峰峰值先除以 2，再映射成单边峰值。",
    )
    add_bullet(
        doc,
        "最终 DAC 输出：I_dac = dc_offset + I_signed，Q_dac = dc_offset + Q_signed，"
        "代码在第 394~409 行。",
    )

    add_text(doc, "3. 各调制模式公式", "Heading 1")
    add_bullet(
        doc,
        "AM：I = A * m * sin(wm t)，Q = 0；代码第 423~431 行。"
        "代码注释明确写了“不额外抬高平均值”。",
    )
    add_bullet(
        doc,
        "FM：phi[n+1] = phi[n] + dphi_fm * sin(wm n)，I = A cos(phi)，Q = A sin(phi)；"
        "实时备用路径在第 434~445 行，主路径走预生成表。",
    )
    add_bullet(
        doc,
        "2ASK：bit=1 时 I=A，bit=0 时 I=0，Q=0；代码第 447~450 行。",
    )
    add_bullet(
        doc,
        "2PSK：bit=1 时 I=+A，bit=0 时 I=-A，Q=0；代码第 453~456 行。"
        "本质上是 BPSK 映射在 I 轴。",
    )
    add_bullet(
        doc,
        "2FSK：连续相位 FSK。bit=1 对应 +shift/2，bit=0 对应 -shift/2，"
        "phi 连续累加，I = A cos(phi)，Q = A sin(phi)；代码第 459~468 行。",
    )

    add_text(doc, "4. 当前实现的重要结论", "Heading 1")
    add_bullet(
        doc,
        "有依据：并不是所有模式都真正使用了旋转的 I/Q 矢量。"
        "AM、2ASK、2PSK 基本只用了 I 路，Q 路为 0。",
    )
    add_bullet(
        doc,
        "有依据：FM、2FSK 才是典型的复基带形式，即 I/Q 两路共同组成相位旋转。"
    )
    add_bullet(
        doc,
        f"代码核实文件：{PROJECT_PATH}",
    )

    doc.save(OUTPUT_PATH)


if __name__ == "__main__":
    build_doc()
